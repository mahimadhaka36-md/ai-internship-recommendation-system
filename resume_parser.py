"""
resume_parser.py
----------------
In-memory resume text extractor supporting PDF and DOCX formats.

Privacy & Security:
- Uploaded resumes are processed purely in-memory (BytesIO) and never stored on disk.
- Resume text is not retained, logged, or shared externally.
- Text-based extraction only; scanned image PDFs trigger a helpful message explaining
  that OCR is not supported in this version.
"""

import io
import os
from typing import Tuple, Optional
import pymupdf
from docx import Document

# Maximum allowed file size in bytes (5 MB)
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024


def parse_resume(file_bytes: bytes, filename: str) -> Tuple[Optional[str], Optional[str]]:
    """
    Extracts text from PDF or DOCX resume bytes in-memory.

    Parameters:
        file_bytes (bytes): Binary content of uploaded file.
        filename (str): Original filename for extension validation.

    Returns:
        (extracted_text, error_message):
            If successful: (text, None)
            If error: (None, error_description)
    """
    if not file_bytes:
        return None, "The uploaded file is empty. Please select a valid resume file."

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        return None, f"File exceeds the 5 MB limit (size: {len(file_bytes) / (1024 * 1024):.2f} MB). Please upload a smaller file."

    ext = os.path.splitext(filename.lower())[1]

    if ext == ".pdf":
        return _extract_from_pdf(file_bytes)
    elif ext in [".docx", ".doc"]:
        if ext == ".doc":
            return None, "Older '.doc' files are not supported. Please save your resume as '.docx' or '.pdf' and re-upload."
        return _extract_from_docx(file_bytes)
    else:
        return None, f"Unsupported file type '{ext}'. Please upload a PDF (.pdf) or Word document (.docx)."


def _extract_from_pdf(file_bytes: bytes) -> Tuple[Optional[str], Optional[str]]:
    """Extract text from PDF bytes using PyMuPDF."""
    try:
        doc = pymupdf.open(stream=file_bytes, filetype="pdf")
        if doc.page_count == 0:
            return None, "The PDF document contains no pages."

        extracted_pages = []
        for page_idx in range(doc.page_count):
            page = doc.load_page(page_idx)
            text = page.get_text()
            if text:
                extracted_pages.append(text.strip())

        full_text = "\n\n".join(extracted_pages).strip()

        if not full_text:
            return (
                None,
                "No readable text could be extracted from this PDF. It appears to be a scanned image or image-only document. "
                "Optical Character Recognition (OCR) is not supported in this initial version. "
                "Please upload a standard text-based PDF or DOCX resume, or enter your skills manually."
            )

        return full_text, None

    except Exception as e:
        return None, f"Unable to parse the PDF document: {str(e)}. Please check that the file is not password-protected or corrupted."


def _extract_from_docx(file_bytes: bytes) -> Tuple[Optional[str], Optional[str]]:
    """Extract text from DOCX bytes using python-docx."""
    try:
        stream = io.BytesIO(file_bytes)
        doc = Document(stream)
        
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract text from tables if any
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        table_texts.append(cell_text)

        full_text = "\n".join(paragraphs + table_texts).strip()

        if not full_text:
            return None, "The DOCX document is empty or contains only non-text media. Please provide a resume with text content."

        return full_text, None

    except Exception as e:
        return None, f"Unable to parse DOCX document: {str(e)}. Please verify the file is a valid .docx document."
